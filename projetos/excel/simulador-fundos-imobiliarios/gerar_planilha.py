import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
perfil_sel = sys.argv[2] if len(sys.argv)>2 else 'Moderado'
wb=Workbook(); ws=wb.active; ws.title='Simulador'; ap=wb.create_sheet('Apoio')
NAVY='0B1F4B'; GRAY='E7E9EE'; INP='FFF4D6'
thin=Side(style='thin',color='BFC5D2'); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
def put(c,v,fmt=None,calc=False,inp=False,bold=False):
    x=ws[c]; x.value=v; x.border=bd
    if fmt: x.number_format=fmt
    if calc: x.fill=PatternFill('solid',fgColor=GRAY)
    if inp: x.fill=PatternFill('solid',fgColor=INP)
    if bold: x.font=Font(bold=True)
    return x
def head(rng,text):
    ws.merge_cells(rng); c=ws[rng.split(':')[0]]; c.value=text
    c.font=Font(bold=True,color='FFFFFF',size=11); c.fill=PatternFill('solid',fgColor=NAVY); c.alignment=Alignment(horizontal='left',indent=1)
ws.sheet_view.showGridLines=False
for col,w in zip('ABCDEFG',[2,34,18,18,18,3,3]): ws.column_dimensions[col].width=w
ws.merge_cells('B2:E3'); b=ws['B2']; b.value='SIMULADOR DE FUNDOS IMOBILIÁRIOS'
b.font=Font(bold=True,size=20,color='FFFFFF'); b.fill=PatternFill('solid',fgColor=NAVY); b.alignment=Alignment(horizontal='center',vertical='center')
ws['B4']='Células amarelas: você edita. Células cinzas: calculadas, não edite. Valores de exemplo, não é recomendação de investimento.'
ws['B4'].font=Font(italic=True,size=9,color='5B6B8C')
# simulação
head('B6:E6','1. Simulação')
put('B7','Aporte mensal (R$)'); put('C7',1500,'R$ #,##0.00',inp=True)
put('B8','Prazo (anos)'); put('C8',10,'0',inp=True)
put('B9','Taxa de rendimento mensal'); put('C9',0.008,'0.00%',inp=True)
put('B10','Patrimônio acumulado (R$)',bold=True); put('C10','=FV(taxa_mensal,prazo_anos*12,-aporte)','R$ #,##0.00',calc=True,bold=True)
put('B11','Dividendos por mês (R$)',bold=True); put('C11','=C10*taxa_mensal','R$ #,##0.00',calc=True,bold=True)
put('B12','Total aportado (R$)'); put('C12','=aporte*prazo_anos*12','R$ #,##0.00',calc=True)
# cenarios
head('B14:E14','2. Cenários (mesmo aporte e taxa)')
for c,t in zip('BCDE',['Prazo (anos)','Meses','Patrimônio (R$)','Dividendos/mês (R$)']):
    put(c+'15',t,bold=True)
for i,a in enumerate([2,5,10,20,30]):
    r=16+i; put(f'B{r}',a,'0'); put(f'C{r}',f'=B{r}*12','0',calc=True)
    put(f'D{r}',f'=FV($C$9,C{r},-$C$7)','R$ #,##0.00',calc=True)
    put(f'E{r}',f'=D{r}*$C$9','R$ #,##0.00',calc=True)
# config
head('B22:E22','3. Configurações')
put('B23','Salário mensal (R$, exemplo)'); put('C23',5000,'R$ #,##0.00',inp=True)
put('B24','% do salário sugerido para aporte'); put('C24',0.3,'0%',inp=True)
put('B25','Aporte sugerido (R$)'); put('C25','=salario*pct_aporte','R$ #,##0.00',calc=True)
put('B26','Rendimento da carteira (a.a. equivalente)'); put('C26','=(1+taxa_mensal)^12-1','0.00%',calc=True)
# perfil
head('B28:E28','4. Divisão do aporte por perfil')
put('B29','Perfil de investidor'); put('C29',perfil_sel,inp=True,bold=True)
dv=DataValidation(type='list',formula1='=lista_perfis',allow_blank=False); ws.add_data_validation(dv); dv.add('C29')
for c,t in zip('BCDE',['Tipo de fundo','Percentual','Valor (R$)','']):
    if t: put(c+'30',t,bold=True)
tipos=['Tijolo (lajes e escritórios)','Logístico','Shoppings','Papel (CRI)','Fundo de fundos','Híbrido']
for i,t in enumerate(tipos):
    r=31+i; put(f'B{r}',t)
    put(f'C{r}',f'=VLOOKUP(perfil&"|"&B{r},tabela_perfis,4,FALSE)','0%',calc=True)
    put(f'D{r}',f'=aporte*C{r}','R$ #,##0.00',calc=True)
put('B37','Total',bold=True); put('C37','=SUM(C31:C36)','0%',calc=True,bold=True); put('D37','=SUM(D31:D36)','R$ #,##0.00',calc=True,bold=True)
# apoio
perf={'Conservador':[10,30,10,30,15,5],'Moderado':[20,20,15,25,10,10],'Arrojado':[25,10,15,40,5,5]}
for c,t in zip('ABCD',['Chave','Perfil','Tipo de fundo','Percentual']):
    ap[c+'1']=t; ap[c+'1'].font=Font(bold=True,color='FFFFFF'); ap[c+'1'].fill=PatternFill('solid',fgColor=NAVY)
r=2
for p,vals in perf.items():
    assert sum(vals)==100
    for t,v in zip(tipos,vals):
        ap[f'A{r}']=f'=B{r}&"|"&C{r}'; ap[f'B{r}']=p; ap[f'C{r}']=t; ap[f'D{r}']=v/100; ap[f'D{r}'].number_format='0%'; r+=1
ap['F1']='Perfis'; ap['F1'].font=Font(bold=True)
for i,p in enumerate(perf): ap[f'F{2+i}']=p
ap['F7']='Os percentuais são exemplos meus (não vêm de nenhuma fonte oficial) e cada perfil soma 100%.'
for c,w in zip('ABCDF',[40,14,30,12,16]): ap.column_dimensions[c].width=w
names={'aporte':"Simulador!$C$7",'prazo_anos':"Simulador!$C$8",'taxa_mensal':"Simulador!$C$9",'salario':"Simulador!$C$23",
 'pct_aporte':"Simulador!$C$24",'perfil':"Simulador!$C$29",'lista_perfis':"Apoio!$F$2:$F$4",'tabela_perfis':f"Apoio!$A$2:$D${r-1}"}
for n,ref in names.items(): wb.defined_names[n]=DefinedName(n,attr_text=ref)
pie=PieChart(); pie.title='Divisão do aporte'; pie.height=7.5; pie.width=11
pie.add_data(Reference(ws,min_col=4,min_row=31,max_row=36)); pie.set_categories(Reference(ws,min_col=2,min_row=31,max_row=36))
pie.dataLabels=DataLabelList(); pie.dataLabels.showPercent=True; pie.dataLabels.showVal=False; pie.dataLabels.showCatName=False; pie.dataLabels.showSerName=False; pie.dataLabels.showLegendKey=False
ws.add_chart(pie,'G6')
from openpyxl.worksheet.properties import PageSetupProperties
ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True)
ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=1
ws.print_area='A1:O38'
wb.save(sys.argv[1])
