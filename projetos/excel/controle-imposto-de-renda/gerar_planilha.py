import sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.properties import PageSetupProperties
NAVY='0B1F4B'; GRAY='E7E9EE'; INP='FFF4D6'; GREEN='1B7F5C'
thin=Side(style='thin',color='BFC5D2'); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
wb=Workbook(); menu=wb.active; menu.title='Menu'
tit=wb.create_sheet('Titular'); inf=wb.create_sheet('Informes'); nts=wb.create_sheet('Notas'); res=wb.create_sheet('Resumo'); ap=wb.create_sheet('Apoio')
def fit(ws):
    ws.sheet_view.showGridLines=False
    ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=1
def banner(ws,rng,text):
    ws.merge_cells(rng); c=ws[rng.split(':')[0]]; c.value=text
    c.font=Font(bold=True,size=18,color='FFFFFF'); c.fill=PatternFill('solid',fgColor=NAVY); c.alignment=Alignment(horizontal='center',vertical='center')
def back(ws,cell):
    ws[cell]='← Voltar ao Menu'; ws[cell].hyperlink="#'Menu'!A1"; ws[cell].font=Font(color='0563C1',underline='single')
def hdr(ws,row,cols,labels):
    for c,l in zip(cols,labels):
        x=ws[f'{c}{row}']; x.value=l; x.font=Font(bold=True,color='FFFFFF'); x.fill=PatternFill('solid',fgColor=NAVY); x.border=bd; x.alignment=Alignment(horizontal='center',wrap_text=True)
def cell(ws,ref,v=None,fmt=None,inp=False,calc=False):
    x=ws[ref]
    if v is not None: x.value=v
    x.border=bd
    if fmt: x.number_format=fmt
    if inp: x.fill=PatternFill('solid',fgColor=INP)
    if calc: x.fill=PatternFill('solid',fgColor=GRAY)
    return x
for w in (menu,tit,inf,nts,res,ap): fit(w)
# Apoio
ap['A1']='Tipo de rendimento'; ap['B1']='Categoria de despesa'; ap['C1']='Tipo de dependente'
for c in 'ABC': ap[f'{c}1'].font=Font(bold=True)
for i,t in enumerate(['Salário','Aluguel','Aposentadoria','Pró-labore','Outros']): ap[f'A{2+i}']=t
for i,t in enumerate(['Saúde','Educação','Previdência privada','Pensão alimentícia','Outras']): ap[f'B{2+i}']=t
for i,t in enumerate(['Nenhum','Filho(a)','Cônjuge','Outro']): ap[f'C{2+i}']=t
for c in 'ABC': ap.column_dimensions[c].width=26
names={'lista_tipos':'Apoio!$A$2:$A$6','lista_categorias':'Apoio!$B$2:$B$6','lista_dependentes':'Apoio!$C$2:$C$5',
 'valor_informes':'Informes!$F$6:$F$105','irrf_informes':'Informes!$G$6:$G$105','tipo_informes':'Informes!$E$6:$E$105',
 'valor_notas':'Notas!$F$6:$F$105','cat_notas':'Notas!$E$6:$E$105'}
for n,r in names.items(): wb.defined_names[n]=DefinedName(n,attr_text=r)
# Menu
banner(menu,'B2:F4','CONTROLE DE DADOS PARA O IMPOSTO DE RENDA')
menu['B6']='Dados de exemplo. Células amarelas: você preenche. Cinzas: calculadas.'; menu['B6'].font=Font(italic=True,size=9,color='5B6B8C')
menu.column_dimensions['A'].width=2
for c in 'BCDEF': menu.column_dimensions[c].width=24
items=[('Titular','Dados do titular e dependentes'),('Informes','Rendimentos e imposto retido na fonte'),('Notas','Despesas dedutíveis (saúde, educação...)'),('Resumo','Totais por tipo e categoria')]
for i,(s,d) in enumerate(items):
    r=8+i*2; menu.merge_cells(f'B{r}:C{r}'); menu.merge_cells(f'D{r}:F{r}')
    b=menu[f'B{r}']; b.value=f'► {s}'; b.hyperlink=f"#'{s}'!A1"; b.font=Font(bold=True,color='FFFFFF',size=12); b.fill=PatternFill('solid',fgColor=GREEN); b.alignment=Alignment(horizontal='center',vertical='center')
    menu[f'D{r}']=d; menu[f'D{r}'].alignment=Alignment(vertical='center',indent=1); menu.row_dimensions[r].height=26
menu['B17']='Links rápidos (abrem no navegador)'; menu['B17'].font=Font(bold=True)
links=[('Receita Federal: Imposto de Renda','https://www.gov.br/receitafederal/pt-br/assuntos/meu-imposto-de-renda'),('Portal gov.br','https://www.gov.br')]
for i,(t,u) in enumerate(links):
    c=menu[f'B{18+i}']; c.value=t; c.hyperlink=u; c.font=Font(color='0563C1',underline='single')
# Titular
banner(tit,'B2:E3','DADOS DO TITULAR'); back(tit,'B4')
tit.column_dimensions['A'].width=2; tit.column_dimensions['B'].width=28; tit.column_dimensions['C'].width=30; tit.column_dimensions['D'].width=22; tit.column_dimensions['E'].width=22
rows=[('Nome completo','Maria Exemplo da Silva',None),('CPF (exemplo fictício)',12345678909,'000"."000"."000"-"00'),('Data de nascimento',datetime.date(1990,5,17),'dd/mm/yyyy'),('E-mail','maria.exemplo@email.com',None),('Ano-calendário',2025,'0')]
for i,(l,v,f) in enumerate(rows):
    r=6+i; cell(tit,f'B{r}',l).font=Font(bold=True); cell(tit,f'C{r}',v,f,inp=True)
hdr(tit,13,'BCD',['Dependente','CPF (exemplo)','Parentesco'])
dep=[('João Exemplo',98765432100,'Filho(a)'),(None,None,None),(None,None,None)]
for i,(n,c,p) in enumerate(dep):
    r=14+i; cell(tit,f'B{r}',n,inp=True); cell(tit,f'C{r}',c,'000"."000"."000"-"00',inp=True); cell(tit,f'D{r}',p,inp=True)
cell(tit,'B18','Nº de dependentes',).font=Font(bold=True); cell(tit,'C18','=COUNTA(B14:B16)','0',calc=True)
d1=DataValidation(type='list',formula1='=lista_dependentes'); tit.add_data_validation(d1); d1.add('D14:D16')
d2=DataValidation(type='custom',formula1='=AND(ISNUMBER(C7),LEN(C7)=11)',showErrorMessage=True,errorTitle='CPF',error='O CPF deve ter 11 dígitos, sem pontos.'); tit.add_data_validation(d2); d2.add('C7')
d3=DataValidation(type='date',operator='lessThan',formula1='TODAY()',showErrorMessage=True,error='Data de nascimento inválida.'); tit.add_data_validation(d3); d3.add('C8')
# Informes
banner(inf,'B2:H3','INFORMES DE RENDIMENTOS'); back(inf,'B4')
for c,w in zip('ABCDEFGH',[2,28,20,14,16,16,16,3]): inf.column_dimensions[c].width=w
hdr(inf,5,'BCDEFG',['Fonte pagadora','CNPJ (exemplo)','Data','Tipo','Valor (R$)','IRRF (R$)'])
ex=[('Empresa Exemplo Ltda',11222333000181,datetime.date(2025,12,31),'Salário',60000,4200),('Imobiliária Exemplo',44555666000199,datetime.date(2025,12,31),'Aluguel',18000,0),('Banco Exemplo S.A.',77888999000122,datetime.date(2025,12,31),'Outros',1200,0)]
for r in range(6,106):
    d=ex[r-6] if r-6<len(ex) else (None,)*6
    cell(inf,f'B{r}',d[0],inp=True); cell(inf,f'C{r}',d[1],'00"."000"."000"/"0000"-"00',inp=True)
    cell(inf,f'D{r}',d[2],'dd/mm/yyyy',inp=True); cell(inf,f'E{r}',d[3],inp=True)
    cell(inf,f'F{r}',d[4],'R$ #,##0.00',inp=True); cell(inf,f'G{r}',d[5],'R$ #,##0.00',inp=True)
v1=DataValidation(type='list',formula1='=lista_tipos',showErrorMessage=True); inf.add_data_validation(v1); v1.add('E6:E105')
v2=DataValidation(type='decimal',operator='greaterThanOrEqual',formula1='0',showErrorMessage=True,error='Informe um valor maior ou igual a zero.'); inf.add_data_validation(v2); v2.add('F6:G105')
inf.freeze_panes='B6'
# Notas
banner(nts,'B2:H3','NOTAS: DESPESAS DEDUTÍVEIS'); back(nts,'B4')
for c,w in zip('ABCDEFGH',[2,28,20,14,20,16,3,3]): nts.column_dimensions[c].width=w
hdr(nts,5,'BCDEF',['Prestador','CNPJ/CPF (exemplo)','Data','Categoria','Valor (R$)'])
nx=[('Clínica Exemplo',12345678000195,datetime.date(2025,3,10),'Saúde',1800),('Escola Exemplo',98765432000110,datetime.date(2025,2,5),'Educação',3500),('Plano de Saúde Exemplo',55443322000166,datetime.date(2025,6,1),'Saúde',4200)]
for r in range(6,106):
    d=nx[r-6] if r-6<len(nx) else (None,)*5
    cell(nts,f'B{r}',d[0],inp=True); cell(nts,f'C{r}',d[1],'0',inp=True); cell(nts,f'D{r}',d[2],'dd/mm/yyyy',inp=True)
    cell(nts,f'E{r}',d[3],inp=True); cell(nts,f'F{r}',d[4],'R$ #,##0.00',inp=True)
v3=DataValidation(type='list',formula1='=lista_categorias',showErrorMessage=True); nts.add_data_validation(v3); v3.add('E6:E105')
v4=DataValidation(type='decimal',operator='greaterThanOrEqual',formula1='0',showErrorMessage=True,error='Informe um valor maior ou igual a zero.'); nts.add_data_validation(v4); v4.add('F6:F105')
nts.freeze_panes='B6'
# Resumo
banner(res,'B2:D3','RESUMO'); back(res,'B4')
for c,w in zip('ABCD',[2,34,20,20]): res.column_dimensions[c].width=w
hdr(res,6,'BC',['Rendimentos por tipo','Total (R$)'])
for i in range(5):
    r=7+i; cell(res,f'B{r}',f'=Apoio!A{2+i}',calc=True); cell(res,f'C{r}',f'=SUMIF(tipo_informes,B{r},valor_informes)','R$ #,##0.00',calc=True)
cell(res,'B12','Total de rendimentos',calc=True).font=Font(bold=True); cell(res,'C12','=SUM(C7:C11)','R$ #,##0.00',calc=True).font=Font(bold=True)
cell(res,'B13','Imposto retido na fonte (IRRF)',calc=True); cell(res,'C13','=SUM(irrf_informes)','R$ #,##0.00',calc=True)
hdr(res,15,'BC',['Despesas por categoria','Total (R$)'])
for i in range(5):
    r=16+i; cell(res,f'B{r}',f'=Apoio!B{2+i}',calc=True); cell(res,f'C{r}',f'=SUMIF(cat_notas,B{r},valor_notas)','R$ #,##0.00',calc=True)
cell(res,'B21','Total de despesas',calc=True).font=Font(bold=True); cell(res,'C21','=SUM(C16:C20)','R$ #,##0.00',calc=True).font=Font(bold=True)
cell(res,'B23','Conferência: informes sem tipo',calc=True); cell(res,'C23','=COUNTA(Informes!B6:B105)-COUNTA(Informes!E6:E105)','0',calc=True)
res['B25']='Os totais não calculam imposto devido: servem para reunir os dados antes de preencher a declaração.'; res['B25'].font=Font(italic=True,size=9,color='5B6B8C')
wb.save('controle-imposto-de-renda.xlsx')
